import os
from io import BytesIO
from datetime import datetime

import streamlit as st
from dotenv import load_dotenv
from openai import OpenAI
from docx import Document
from docx.shared import Pt

load_dotenv()

# 1) PASSWORD PROTECTION
def get_password():
    """
    เช็ค PASSWORD จาก st.secrets ก่อน (สำหรับ Streamlit Cloud)
    แล้ว fallback ไปที่ .env / environment variable (สำหรับรัน local)
    ห่อด้วย try/except กัน crash ตอนไม่มีไฟล์ secrets.toml เลย
    เหมือนกับที่ทำไว้ใน get_api_key()
    """
    try:
        if "PASSWORD" in st.secrets:
            return st.secrets["PASSWORD"]
    except Exception:
        pass
    return os.environ.get("PASSWORD")


APP_PASSWORD = get_password()

password = st.text_input(
    "Enter password",
    type="password"
)

if not APP_PASSWORD:
    st.error(
        "ยังไม่ได้ตั้งค่า PASSWORD — "
        "ถ้ารัน local ให้ใส่ PASSWORD=... ใน .env, "
        "ถ้า deploy บน Streamlit Cloud ให้ใส่ใน App settings > Secrets"
    )
    st.stop()

if password != APP_PASSWORD:
    if password:
        st.error("Incorrect password")
    else:
        st.info("Please enter the password.")

    st.stop()

# 2) APP CONTENT
st.success("Access granted")
# 3) ตั้งค่า OpenAI client
# รองรับทั้งการรัน local (env var)
# และบน Streamlit Cloud (st.secrets)

def get_api_key():
    try:
        if "OPENAI_API_KEY" in st.secrets:
            return st.secrets["OPENAI_API_KEY"]
    except Exception:
        pass

    return os.environ.get("OPENAI_API_KEY")


api_key = get_api_key()

client = OpenAI(api_key=api_key) if api_key else None


ALLOWED_MODELS = [
    "gpt-5-pro",
    "gpt-4o",
    "gpt-4o-mini"
]

DEFAULT_MODEL = "gpt-4o"
# 4) CONTRACT KNOWLEDGE
CONTRACT_ELEMENTS_KNOWLEDGE = """
คุณคือผู้ช่วยกฎหมาย (legal assistant) ที่เชี่ยวชาญเรื่อง "การเกิดขึ้นของสัญญา" (Contract Formation) ภายใต้กฎหมายทั่วไปของออสเตรเลีย (โดยเฉพาะรัฐวิกตอเรีย)อ้างอิงจาก The Law Handbook บทที่ 7.1 "Elements of a Contract"

ซึ่งมีองค์ประกอบหลัก 6 ข้อดังนี้:
1) Offer and Acceptance (คำเสนอและคำสนอง)
   - สัญญาเกิดขึ้นเมื่อคำเสนอ (offer) ของฝ่ายหนึ่งได้รับการสนอง (accepted) จากอีกฝ่ายหนึ่ง
   - คำเสนอต้องเป็นคำมั่นที่ชัดเจนและแน่นอนว่าจะผูกพันไม่ใช่เพียงการแสดงความเต็มใจจะเจรจาต่อ
   - การสนองต้องตรงกับสิ่งที่เสนอเป๊ะๆ (mirror image)
   - หากเปลี่ยนแปลงเงื่อนไข จะถือเป็น "คำเสนอโต้กลับ"(counter-offer)
   - ผู้เสนอสามารถถอนคำเสนอได้ก่อนมีการสนองรับแต่ต้องสื่อสารการถอนนั้นไปยังอีกฝ่ายอย่างชัดเจน
   - การนิ่งเฉยไม่ถือเป็นการยอมรับ
   - การสนองรับอาจเกิดจากพฤติกรรม (conduct) ได้เช่นกัน

2) Intention to Create Legal Relations
   (เจตนาที่จะผูกพันทางกฎหมาย)
   - แค่มีข้อตกลงกันยังไม่พอ
   - คู่สัญญาต้องมีเจตนาที่จะให้ข้อตกลงนั้นมีผลผูกพันทางกฎหมาย
   - พิจารณาจากสิ่งที่ "บุคคลที่มีเหตุผล" (reasonable person)
     จะเข้าใจจากสถานการณ์
   - ในการทำธุรกิจ/การค้าทั่วไป (commercial/arm's length)
     กฎหมายจะสันนิษฐานไว้ก่อนว่ามีเจตนาผูกพันทางกฎหมาย
   - ในความสัมพันธ์ทางครอบครัวหรือเพื่อนฝูง/ข้อตกลงเชิงสังคม
     มักสันนิษฐานว่าไม่มีเจตนาผูกพันทางกฎหมาย

3) Consideration (ค่าตอบแทน/สิ่งตอบแทน)
   - คือ "ราคา" ที่จ่ายเพื่อแลกกับคำมั่นของอีกฝ่าย
   - ไม่จำเป็นต้องเป็นเงิน
   - ต้องมีมูลค่าบางอย่าง
   - ศาลจะไม่ตัดสินความเหมาะสมของมูลค่า
   - ความรัก ความเสน่หา หรือของกำนัลโดยสมัครใจ
     ไม่ถือเป็น consideration ที่สมบูรณ์
   - Deed ไม่จำเป็นต้องมี consideration

4) Legal Capacity (ความสามารถตามกฎหมายในการทำสัญญา)
   - บุคคลบางกลุ่มอาจมีข้อจำกัดในการทำสัญญาผูกพัน
   - ผู้เยาว์ ผู้มีความบกพร่องทางจิต
     บุคคลล้มละลาย นิติบุคคล และผู้ต้องขัง
     อาจมีข้อจำกัดแตกต่างกัน
   - ผู้เยาว์อาจผูกพันในสัญญาเกี่ยวกับ necessaries
   - ผู้มีความบกพร่องทางจิตอาจไม่ผูกพัน
     หากไม่เข้าใจลักษณะทั่วไปของสัญญา
     และอีกฝ่ายรู้หรือควรรู้ถึงความบกพร่องนั้น

5) Consent (ความยินยอมโดยแท้จริง)
   - การทำสัญญาต้องเกิดจากเจตจำนงเสรี
   - ความยินยอมอาจถูกกระทบโดย:
     mistake, misrepresentation, duress, undue influence, unconscionability หรือ unfair terms
   - หากความยินยอมถูกกระทบอย่างมีนัยสำคัญ
     สัญญาอาจเป็นโมฆียะ (voidable)

6) Legality (ความชอบด้วยกฎหมายของสัญญา)

   - สัญญาที่มีวัตถุประสงค์ผิดกฎหมายหรือขัดต่อนโยบายสาธารณะจะไม่สามารถบังคับได้ตามกฎหมาย

หลักการวิเคราะห์
----------------
เมื่อผู้ใช้ป้อนเหตุการณ์เข้ามา ให้คุณ:
1. วิเคราะห์เหตุการณ์นั้นเทียบกับองค์ประกอบทั้ง 6 ข้อ
   ทีละข้อ
2. สำหรับแต่ละองค์ประกอบที่เกี่ยวข้อง
   ให้อธิบายหลักกฎหมายสั้นๆ ก่อน
   แล้วจึงปรับใช้กับข้อเท็จจริง
3. สรุปในตอนท้ายอย่างชัดเจนว่า
   "สัญญาเกิดขึ้นแล้วหรือยังไม่เกิดขึ้น"
   Contract formed /
   Not yet formed /
   Uncertain - need more facts
4. หากข้อมูลไม่เพียงพอให้บอกว่าต้องการข้อเท็จจริงเพิ่มเติมอะไร
5. ตอบเป็นภาษาเดียวกับที่ผู้ใช้พิมพ์เข้ามา
6. คำตอบต้องอิงตามหลักกฎหมายของออสเตรเลีย (Victoria) ตามที่สรุปไว้ข้างต้นเท่านั้น
7. ระบุด้วยว่านี่เป็นข้อมูลเพื่อการศึกษาเบื้องต้นเท่านั้นไม่ใช่คำปรึกษาทางกฎหมายอย่างเป็นทางการ

หากผู้ใช้ถามคำถามต่อเนื่อง (follow-up question) หลังจากคุณวิเคราะห์เหตุการณ์แรกไปแล้ว ให้ตอบโดยอ้างอิงกับเหตุการณ์และบทวิเคราะห์เดิม
ในบทสนทนานี้ด้วย
"""

SYSTEM_PROMPT = CONTRACT_ELEMENTS_KNOWLEDGE
# ============================================================
# 5) OPENAI API
# ============================================================
def call_openai(history: list, model_name: str) -> str:

    messages = [
        {
            "role": "system",
            "content": SYSTEM_PROMPT
        }
    ] + history

    response = client.chat.completions.create(
        model=model_name,
        messages=messages,
        temperature=0.2,
    )

    return response.choices[0].message.content


def build_initial_prompt(scenario_text: str) -> str:

    return (
        "Please analyze the situation based on the 6 elements "
        "of contract formation and summarize:\n\n"
        f"situation: {scenario_text}"
    )

# ============================================================
# 6) BUILD DOCX
# ============================================================
def build_docx(messages: list) -> BytesIO:

    doc = Document()

    title = doc.add_heading(
        "Contract Formation Chatbot — Discussion Log",
        level=1
    )

    subtitle = doc.add_paragraph(
        f"Exported on: "
        f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )

    subtitle.runs[0].italic = True

    doc.add_paragraph(
        "Reference: The Law Handbook (Victoria, Australia) — "
        "Elements of a Contract"
    )

    doc.add_paragraph()

    role_labels = {
        "user": "You",
        "assistant": "Assistant"
    }

    for msg in messages:

        role = role_labels.get(
            msg["role"],
            msg["role"].capitalize()
        )

        heading = doc.add_paragraph()

        run = heading.add_run(f"{role}:")
        run.bold = True
        run.font.size = Pt(12)

        for line in msg["content"].split("\n"):
            doc.add_paragraph(
                line if line.strip() else ""
            )

        doc.add_paragraph()

    buffer = BytesIO()

    doc.save(buffer)

    buffer.seek(0)

    return buffer

# ============================================================
# 7) STREAMLIT PAGE
# ============================================================
st.set_page_config(
    page_title="Contract Formation Chatbot (Australia)",
    page_icon="⚖️"
)

st.title("⚖️ Contract Formation Chatbot (Australia)")

st.caption(
    "This chatbot will answer your question about "
    "the contract formation in Victoria, Australia only, "
    "using the reference from The Law Handbook — "
    "Elements of a Contract."
)
# 8) CHECK OPENAI API KEY
if not api_key:

    st.error(
        "ไม่พบ OPENAI_API_KEY — "
        "ถ้ารัน local ให้ตั้งค่า environment variable, "
        "ถ้า deploy บน Streamlit Cloud "
        "ให้ใส่ใน App settings > Secrets"
    )
# 9) SIDEBAR
with st.sidebar:

    st.subheader("Settings")

    model_name = st.selectbox(
        "Chatbot Models",
        ALLOWED_MODELS,
        index=ALLOWED_MODELS.index(DEFAULT_MODEL)
    )

    if st.button("Clear discussion history"):

        st.session_state.messages = []
        st.session_state.scenario_submitted = False
# 10) SESSION STATE
if "messages" not in st.session_state:

    st.session_state.messages = []


if "scenario_submitted" not in st.session_state:

    st.session_state.scenario_submitted = False
# 11) SHOW CHAT HISTORY
for msg in st.session_state.messages:

    with st.chat_message(msg["role"]):

        st.markdown(msg["content"])
# 12) INITIAL SCENARIO
if not st.session_state.scenario_submitted:

    with st.form(
        key="scenario_form",
        clear_on_submit=True
    ):

        scenario_text = st.text_area(
            "Please provide your situation here.",
            placeholder=(
                "i.e. 'I offered him $500 for this camera "
                "but he did not agree.'"
            ),
            height=120,
        )

        submitted = st.form_submit_button("Submit")


    if submitted:

        if not scenario_text or not scenario_text.strip():

            st.warning(
                "Please provide the situation before submit."
            )

        elif not api_key:

            st.warning(
                "ยังไม่ได้ตั้งค่า OPENAI_API_KEY "
                "จึงเรียก AI ไม่ได้"
            )

        else:

            user_prompt = build_initial_prompt(
                scenario_text
            )

            st.session_state.messages.append(
                {
                    "role": "user",
                    "content": user_prompt
                }
            )

            with st.chat_message("user"):

                st.markdown(scenario_text)


            with st.chat_message("assistant"):

                with st.spinner("assessing..."):

                    try:

                        answer = call_openai(
                            st.session_state.messages,
                            model_name
                        )

                    except Exception as e:

                        answer = (
                            "เกิดข้อผิดพลาดในการเรียก "
                            f"OpenAI API: {e}"
                        )

                    st.markdown(answer)


            st.session_state.messages.append(
                {
                    "role": "assistant",
                    "content": answer
                }
            )

            st.session_state.scenario_submitted = True

            st.rerun()
# 13) FOLLOW-UP QUESTIONS
if st.session_state.scenario_submitted:

    followup_text = st.chat_input(
        "Any follow-up question?"
    )


    if followup_text:

        if not api_key:

            st.warning(
                "ยังไม่ได้ตั้งค่า OPENAI_API_KEY "
                "จึงเรียก AI ไม่ได้"
            )

        else:

            st.session_state.messages.append(
                {
                    "role": "user",
                    "content": followup_text
                }
            )


            with st.chat_message("user"):

                st.markdown(followup_text)


            with st.chat_message("assistant"):

                with st.spinner("assessing..."):

                    try:

                        answer = call_openai(
                            st.session_state.messages,
                            model_name
                        )

                    except Exception as e:

                        answer = (
                            "เกิดข้อผิดพลาดในการเรียก "
                            f"OpenAI API: {e}"
                        )

                    st.markdown(answer)


            st.session_state.messages.append(
                {
                    "role": "assistant",
                    "content": answer
                }
            )


    if st.button("Analyze new question."):

        st.session_state.messages = []

        st.session_state.scenario_submitted = False

        st.rerun()
# 14) DOWNLOAD DISCUSSION
if st.session_state.messages:

    st.divider()

    docx_buffer = build_docx(
        st.session_state.messages
    )

    st.download_button(
        label="⬇️ Download this discussion",
        data=docx_buffer,
        file_name=(
            "contract_formation_discussion_"
            f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        ),
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
    )